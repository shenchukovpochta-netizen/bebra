"""Инструменты списков: сортировка, страница, итоги, свои фильтры, xlsx.

Один набор правил на все списки: парк, аренды, наряды и склад со своим
устройством разъедутся на первой же правке.

Два правила, которые здесь стерегут. Первое: сортировать можно только по
колонкам из белого списка — имя поля из адреса это чужая строка. Второе:
итог в подвале считается по всему найденному, а не по видимой странице:
«итого 77 аренд» при пятидесяти на экране — это и есть ответ на вопрос,
ради которого список открывали.
"""

from __future__ import annotations

import io
import sys
import unittest
import zipfile
from decimal import Decimal
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from app.crm import logic  # noqa: E402

try:
    import test_web as tw
    HAVE_WEB = tw.HAVE_WEB
except ImportError:                                    # pragma: no cover
    HAVE_WEB = False

D = Decimal


def _run(coro):
    import asyncio
    return asyncio.run(coro)


class TestSortRows(unittest.TestCase):
    def rows(self):
        return [{"code": "B-2", "mileage_km": 500, "full_name": "Яшин"},
                {"code": "B-1", "mileage_km": 1500, "full_name": None},
                {"code": "B-3", "mileage_km": None, "full_name": "Аверин"}]

    def allowed(self):
        return {"code": "code", "mileage": "mileage_km", "client": "full_name"}

    def test_text_sorting(self):
        got = logic.sort_rows(self.rows(), "code", "asc", allowed=self.allowed())
        self.assertEqual([r["code"] for r in got], ["B-1", "B-2", "B-3"])

    def test_reverse(self):
        got = logic.sort_rows(self.rows(), "code", "desc", allowed=self.allowed())
        self.assertEqual([r["code"] for r in got], ["B-3", "B-2", "B-1"])

    def test_numbers_are_not_strings(self):
        got = logic.sort_rows(self.rows(), "mileage", "asc", allowed=self.allowed())
        self.assertEqual([r["mileage_km"] for r in got][:2], [500, 1500],
                         "иначе 1500 встало бы перед 500")

    def test_empty_goes_last_both_ways(self):
        up = logic.sort_rows(self.rows(), "client", "asc", allowed=self.allowed())
        down = logic.sort_rows(self.rows(), "client", "desc", allowed=self.allowed())
        self.assertIsNone(up[-1]["full_name"])
        self.assertIsNone(down[-1]["full_name"],
                          "«сортировка по клиенту» не должна выносить наверх "
                          "всё, что ещё не выдано")

    def test_unknown_column_changes_nothing(self):
        got = logic.sort_rows(self.rows(), "password", "asc", allowed=self.allowed())
        self.assertEqual([r["code"] for r in got], ["B-2", "B-1", "B-3"])

    def test_no_column_changes_nothing(self):
        got = logic.sort_rows(self.rows(), "", "asc", allowed=self.allowed())
        self.assertEqual([r["code"] for r in got], ["B-2", "B-1", "B-3"])


class TestPageOf(unittest.TestCase):
    def rows(self, n=120):
        return [{"i": i} for i in range(n)]

    def test_first_page(self):
        page = logic.page_of(self.rows(), 50, 1)
        self.assertEqual(len(page["rows"]), 50)
        self.assertEqual(page["total"], 120, "итог - по всему найденному")
        self.assertEqual(page["pages"], 3)
        self.assertTrue(page["has_more"])

    def test_last_page_is_shorter(self):
        page = logic.page_of(self.rows(), 50, 3)
        self.assertEqual(page["shown"], 20)
        self.assertEqual(page["rows"][0]["i"], 100)

    def test_page_out_of_range_is_clamped(self):
        self.assertEqual(logic.page_of(self.rows(), 50, 99)["page"], 3)
        self.assertEqual(logic.page_of(self.rows(), 50, 0)["page"], 1)
        self.assertEqual(logic.page_of(self.rows(), 50, "ерунда")["page"], 1)

    def test_size_outside_the_list_falls_back(self):
        self.assertEqual(logic.page_of(self.rows(), 7, 1)["size"],
                         logic.DEFAULT_LIST_SIZE)
        self.assertEqual(logic.check_list_size("300"), 300)
        self.assertEqual(logic.check_list_size("1000000"),
                         logic.DEFAULT_LIST_SIZE,
                         "«покажи всё» упирается в 300: больше не читает никто")

    def test_empty_list(self):
        page = logic.page_of([], 50, 1)
        self.assertEqual(page["total"], 0)
        self.assertEqual(page["pages"], 1)
        self.assertFalse(page["has_more"])

    def test_sum_of_column(self):
        self.assertEqual(logic.sum_of([{"x": D(100)}, {"x": D("50.5")},
                                       {"x": None}], "x"), D("150.50"))


@unittest.skipUnless(HAVE_WEB, "нет fastapi/httpx")
class TestListsInThePanel(tw.WebCase):
    def setUp(self):
        super().setUp()
        self.login()
        self.seed()
        for i in range(2, 8):
            _run(self.crm.create_bike(code=f"B-{i}", model="Kugoo V3",
                                      mileage_km=i * 100))

    def test_sorting_reorders_the_table(self):
        text = self.get_ok("/bikes?sort=code&dir=desc")
        first = text.index("B-7")
        self.assertLess(first, text.index("B-1"), "по убыванию номера")

    def test_sort_link_flips_the_direction(self):
        text = self.get_ok("/bikes?sort=code&dir=asc")
        self.assertIn("sort=code&dir=desc", text,
                      "повторный клик переворачивает порядок")

    def test_unknown_sort_is_ignored(self):
        self.assertEqual(self.client.get("/bikes?sort=password").status_code, 200)

    def test_page_size_and_total(self):
        text = self.get_ok("/bikes?rows=50")
        self.assertIn("Итого 7", text)
        self.assertIn("строк:", text)

    def test_one_page_has_no_pager(self):
        # Со страницей в 50 строк семь велосипедов умещаются на одной,
        # и листалка там только мешает.
        self.assertNotIn("стр. 1 из", self.get_ok("/bikes?rows=50"))

    def test_lists_carry_the_tools(self):
        for url in ("/bikes", "/rentals", "/orders", "/parts"):
            with self.subTest(url=url):
                text = self.get_ok(url)
                self.assertIn("Итого", text)
                self.assertIn("Мои фильтры", text)


@unittest.skipUnless(HAVE_WEB, "нет fastapi/httpx")
class TestSavedViews(tw.WebCase):
    def setUp(self):
        super().setUp()
        self.login()
        self.seed()

    def save(self, name="Мои должники", query="status=active&view=debt"):
        return self.client.post("/views", data={
            "section": "/rentals", "name": name, "query": query})

    def test_save_and_show(self):
        self.save()
        text = self.get_ok("/rentals")
        self.assertIn("Мои должники", text)
        self.assertIn("view=debt", text)

    def test_same_name_overwrites(self):
        self.save(query="status=active")
        self.save(query="status=closed")
        staff = _run(self.crm.staff_by_login("admin"))
        views = _run(self.crm.saved_views(staff["id"], "/rentals"))
        self.assertEqual(len(views), 1, "второе сохранение обновляет набор")
        self.assertEqual(views[0]["query"], "status=closed")

    def test_delete_own(self):
        self.save()
        staff = _run(self.crm.staff_by_login("admin"))
        view = _run(self.crm.saved_views(staff["id"], "/rentals"))[0]
        self.client.post(f"/views/{view['id']}/delete")
        self.assertEqual(_run(self.crm.saved_views(staff["id"], "/rentals")), [])

    def test_someone_elses_view_is_not_touched(self):
        self.save()
        staff = _run(self.crm.staff_by_login("admin"))
        view = _run(self.crm.saved_views(staff["id"], "/rentals"))[0]
        boss = _run(self.crm.staff_by_login("admin"))
        _run(self.crm.create_staff("drugoy", logic.hash_password("drugoy-pass"),
                                   name="Другой", role="admin",
                                   profile_id=boss.get("profile_id")))
        self.client.post("/logout")
        self.login("drugoy", "drugoy-pass")
        self.client.post(f"/views/{view['id']}/delete")
        self.assertEqual(len(_run(self.crm.saved_views(staff["id"], "/rentals"))), 1,
                         "id в адресе - чужая строка, и владелец проверяется")

    def test_views_are_personal(self):
        self.save()
        staff = _run(self.crm.staff_by_login("admin"))
        _run(self.crm.create_staff("sosed", logic.hash_password("sosed-pass"),
                                   name="Сосед", role="admin",
                                   profile_id=staff.get("profile_id")))
        self.client.post("/logout")
        self.login("sosed", "sosed-pass")
        self.assertNotIn("Мои должники", self.get_ok("/rentals"))


@unittest.skipUnless(HAVE_WEB, "нет fastapi/httpx")
class TestExcelExport(tw.WebCase):
    def setUp(self):
        super().setUp()
        self.login()
        self.seed()

    def book(self, url):
        r = self.client.get(url)
        self.assertEqual(r.status_code, 200, url)
        self.assertIn("spreadsheetml", r.headers["content-type"])
        self.assertTrue(zipfile.is_zipfile(io.BytesIO(r.content)),
                        "xlsx - это zip")
        import openpyxl
        return openpyxl.load_workbook(io.BytesIO(r.content))

    def test_bikes_xlsx(self):
        sheet = self.book("/bikes.xlsx").active
        self.assertEqual(sheet.cell(row=1, column=1).value, "Номер")
        self.assertEqual(sheet.cell(row=2, column=1).value, "B-1")
        self.assertEqual(sheet.freeze_panes, "A2", "шапка не уезжает")

    def test_numbers_stay_numbers(self):
        _run(self.crm.update_bike(self.bike_id, mileage_km=1234, by="тест"))
        sheet = self.book("/bikes.xlsx").active
        header = [c.value for c in sheet[1]]
        col = header.index("Пробег, км") + 1
        self.assertEqual(sheet.cell(row=2, column=col).value, 1234)
        self.assertIsInstance(sheet.cell(row=2, column=col).value, int,
                              "иначе сумму в Excel не поставить")

    def test_formula_looking_text_is_not_a_formula(self):
        _run(self.crm.create_bike(code="=HYPERLINK(1)", model="Kugoo V3"))
        sheet = self.book("/bikes.xlsx").active
        values = [sheet.cell(row=i, column=1).value for i in range(2, 4)]
        self.assertIn("=HYPERLINK(1)", values)
        for cell in sheet["A"]:
            self.assertNotEqual(cell.data_type, "f", "формулой это не стало")

    def test_csv_still_works(self):
        r = self.client.get("/bikes.csv")
        self.assertEqual(r.status_code, 200)
        self.assertIn("text/csv", r.headers["content-type"])

    def test_unknown_format_is_404(self):
        self.assertEqual(self.client.get("/bikes.pdf").status_code, 404,
                         "молча отдать csv значило бы соврать в имени файла")

    def test_every_export_has_both(self):
        for stem in ("/bikes", "/rentals", "/orders", "/parts"):
            with self.subTest(stem=stem):
                self.assertEqual(self.client.get(stem + ".xlsx").status_code, 200)
                self.assertEqual(self.client.get(stem + ".csv").status_code, 200)


if __name__ == "__main__":                              # pragma: no cover
    unittest.main()
