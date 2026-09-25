"""Демо-стенд: режим панели (плашка, входы под ролями, запреты, noindex,
503 на время сброса) и процесс app.demo (конфиг без боевых ключей, ночной
сброс). Панель - на заглушке FakeCrm, без базы и без uvicorn.
"""

from __future__ import annotations

import asyncio
import io
import os
import re
import sys
import tempfile
import unittest
from datetime import UTC, date, datetime, timedelta, timezone
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

try:
    from fastapi.testclient import TestClient

    from app.crm import logic, service
    from app.web import app as web_app
    from app.web.app import create_app
    from app.web.config import WebConfig
    from tests.fake_crm import FakeCrm
    from tests.test_web import FakeBotDB, run
    HAVE_WEB = True
except ImportError:                                    # pragma: no cover
    HAVE_WEB = False

try:
    from app.demo import runtime, seed, seed_extras
    from app.services.crypto import Vault
    HAVE_DEMO = HAVE_WEB
except ImportError:                                    # pragma: no cover
    HAVE_DEMO = False

BANNER = "Демо-версия:"
BLOCKED = "В демо-версии это недоступно"
MSK = timezone(timedelta(hours=3), "MSK")
PROFILE_OF = {"demo": "owner", "operator": "manager", "mechanic": "tech"}


@unittest.skipUnless(HAVE_WEB, "fastapi не установлен")
class DemoCase(unittest.TestCase):
    """Панель в режиме демо с тремя логинами сида (пароль demo)."""

    demo = True

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.crm = FakeCrm()
        self.cfg = WebConfig(pg={}, secret="test-secret", admin_login="admin",
                             admin_password="", bot_token="",
                             storage_dir=Path(self.tmp.name) / "kyc", port=8080,
                             remind_before_days=2, demo=self.demo,
                             bike_photo_dir=Path(self.tmp.name) / "bikes",
                             doc_dir=Path(self.tmp.name) / "doctemplates")
        self.staff_ids = {}
        for login, code in PROFILE_OF.items():
            profile = run(self.crm.access_profile_by_code(code))
            self.staff_ids[login] = run(self.crm.create_staff(
                login, logic.hash_password("demo"), login.title(),
                "admin" if code == "owner" else "manager", profile["id"]))
        self.app = create_app(crm=self.crm, db=FakeBotDB(), cfg=self.cfg, bot=None)
        self.client = self.client_from("testclient")

    def client_from(self, host):
        return TestClient(self.app, follow_redirects=False, client=(host, 50000))

    def login(self, login="demo", password="demo", client=None):
        return (client or self.client).post("/login",
                                            data={"login": login, "password": password})

    def get_ok(self, path):
        r = self.client.get(path)
        self.assertEqual(r.status_code, 200, f"{path}: {r.status_code}")
        return r.text

    def hash_of(self, login):
        return self.crm.staff[self.staff_ids[login]]["password_hash"]


class TestBannerAndLogins(DemoCase):
    def test_banner_on_login_pages_and_denied(self):
        self.assertIn(BANNER, self.get_ok("/login"))
        self.assertEqual(self.login().status_code, 303)
        for path in ("/", "/bikes", "/me"):
            self.assertIn(BANNER, self.get_ok(path), path)
        # отказ по правам - тоже страница панели, и плашка на ней есть
        mechanic = self.client_from("10.0.0.2")
        self.assertEqual(self.login("mechanic", client=mechanic).status_code, 303)
        r = mechanic.get("/staff")
        self.assertEqual(r.status_code, 403)
        self.assertIn(BANNER, r.text)
        # страница подписи клиента - свой шаблон, плашка и там
        r = self.client.get("/sign/" + "x" * 32)
        self.assertEqual(r.status_code, 404)
        self.assertIn(BANNER, r.text)

    def test_login_page_offers_three_roles_without_script(self):
        text = self.get_ok("/login")
        forms = re.findall(r'<form method="post" action="/login">(.*?)</form>', text, re.S)
        pairs = set()
        for body in forms:
            login = re.search(r'name="login" value="([^"]+)"', body)
            password = re.search(r'name="password" value="([^"]+)"', body)
            if login and password:
                pairs.add((login.group(1), password.group(1)))
        self.assertEqual(pairs, {("demo", "demo"), ("operator", "demo"),
                                 ("mechanic", "demo")})
        self.assertNotIn("<script", text)
        for login, _, role in web_app.DEMO_LOGINS:
            self.assertIn(role, text)
            self.assertIn(f"<code>{login}</code>", text)

    def test_every_hint_logs_in(self):
        for i, (login, password, _) in enumerate(web_app.DEMO_LOGINS):
            client = self.client_from(f"10.0.1.{i}")
            r = self.login(login, password, client=client)
            self.assertEqual(r.status_code, 303, login)

    @unittest.skipUnless(HAVE_DEMO, "нет пакета демо")
    def test_hints_match_the_seed(self):
        """Панель пакет демо не импортирует - список на входе сверяется здесь."""
        seeded = {login: profile for login, _, profile, _, _, shown in seed.STAFF if shown}
        self.assertEqual({login for login, _, _ in web_app.DEMO_LOGINS}, set(seeded))
        for login, password, _ in web_app.DEMO_LOGINS:
            self.assertEqual(password, "demo")
            self.assertEqual(seeded[login], PROFILE_OF[login])

    def test_session_cookie_is_its_own(self):
        """Боевая панель и демо на соседних портах localhost не выбивают
        друг друга: cookie портов не различают."""
        r = self.login()
        self.assertIn("crm_demo=", r.headers.get("set-cookie", ""))


class TestBlockedWrites(DemoCase):
    def setUp(self):
        super().setUp()
        self.assertEqual(self.login().status_code, 303)
        self.victim = self.staff_ids["operator"]

    def post(self, path, **kwargs):
        r = self.client.post(path, headers={"referer": "http://testserver/me"}, **kwargs)
        self.assertEqual(r.status_code, 303, path)
        self.assertEqual(r.headers["location"], "/me", path)
        return r

    def test_every_blocked_route_writes_nothing(self):
        crm = self.crm
        staff_before = {k: dict(v) for k, v in crm.staff.items()}
        profiles_before = {k: dict(v) for k, v in crm.profiles_.items()}
        settings_before = dict(crm.settings_)
        manager = run(crm.access_profile_by_code("manager"))
        vid = self.victim
        routes = [
            ("/me/password", {"data": {"old": "demo", "new": "new-pass-123",
                                       "new2": "new-pass-123"}}),
            ("/staff", {"data": {"login": "intruder", "password": "intruder-123",
                                 "name": "Чужой", "profile_id": str(manager["id"])}}),
            (f"/staff/{vid}/profile", {"data": {"profile_id": str(manager["id"])}}),
            (f"/staff/{vid}/location", {"data": {"location": "Павлюхина"}}),
            (f"/staff/{vid}/telegram", {"data": {"action": "code"}}),
            (f"/staff/{vid}/password", {"data": {"password": "reset-pass-123"}}),
            (f"/staff/{vid}/toggle", {}),
            ("/profiles", {"data": {"name": "Свой", "sec_bikes": "edit"}}),
            (f"/profiles/{manager['id']}", {"data": {"name": "Переименован"}}),
            (f"/profiles/{manager['id']}/delete", {}),
            ("/documents/contract", {"files": {"template": ("мой.docx", b"PK\x03\x04",
                                                            "application/octet-stream")}}),
            ("/documents/contract", {"data": {"action": "ours"}}),
            ("/documents/marks/stamp", {"files": {"mark": ("печать.png", b"\x89PNG",
                                                           "image/png")}}),
            ("/payments/acquiring", {"data": {"action": "on"}}),
            ("/payments/acquiring", {"data": {"action": "check"}}),
        ]
        for path, kwargs in routes:
            self.post(path, **kwargs)
        self.assertEqual({k: dict(v) for k, v in crm.staff.items()}, staff_before)
        self.assertEqual({k: dict(v) for k, v in crm.profiles_.items()}, profiles_before)
        self.assertEqual(crm.settings_, settings_before)
        self.assertEqual(run(crm.doc_templates()), [])
        self.assertEqual(run(crm.company_marks()), [])
        self.assertFalse(Path(self.cfg.doc_dir).exists())
        # после отказа - на ту же страницу, с объяснением
        self.assertIn(BLOCKED, self.get_ok("/me"))
        # и сессия жива: пароль не менялся, входить заново не надо
        self.assertTrue(logic.verify_password("demo", self.hash_of("demo")))

    def test_back_only_to_own_pages(self):
        r = self.client.post("/me/password", data={"old": "demo"},
                             headers={"referer": "https://evil.example/staff"})
        self.assertEqual(r.headers["location"], "/")
        r = self.client.post("/me/password", data={"old": "demo"})
        self.assertEqual(r.headers["location"], "/")
        r = self.client.post("/staff/1/toggle",
                             headers={"referer": "http://testserver/staff?q=1"})
        self.assertEqual(r.headers["location"], "/staff?q=1")

    def test_reading_stays_open(self):
        for path in ("/staff", "/profiles", "/documents", "/payments", "/me"):
            self.get_ok(path)

    def test_list_is_exact(self):
        blocked = web_app.demo_blocked
        for path in ("/me/password", "/payments/acquiring", "/staff", "/staff/7/password",
                     "/profiles", "/profiles/3/delete", "/documents/contract",
                     "/documents/marks/stamp", "/import"):
            self.assertTrue(blocked("POST", path), path)
            self.assertFalse(blocked("GET", path), path)
        for path in ("/staffing", "/documentsx", "/payments", "/payments/5",
                     "/bikes/1/check", "/me", "/login", "/rentals"):
            self.assertFalse(blocked("POST", path), path)

    def test_import_is_closed(self):
        """Импорт таблицы: тысячи строк одним запросом перекосили бы три
        числа до ночи, а чужая таблица с настоящими клиентами стала бы
        видна всем. Разбор файла не начинается вовсе - ни сухой, ни с
        записью (разбор xlsx - и есть самый тяжёлый запрос)."""
        from tests.test_import import ROWS, sheet
        workbook = sheet(ROWS)
        with mock.patch.object(web_app.import_xlsx, "run") as run_import:
            for apply in ("0", "1"):
                r = self.client.post("/import", data={"apply": apply},
                                     files={"file": ("учёт.xlsx", workbook,
                                                     "application/octet-stream")},
                                     headers={"referer": "http://testserver/import"})
                self.assertEqual((r.status_code, r.headers["location"]), (303, "/import"))
            run_import.assert_not_called()
        text = self.get_ok("/import")
        self.assertIn(BLOCKED, text)
        self.assertIn("импорт выключен", text)
        self.assertNotIn('action="/import"', text)

    def test_other_writes_still_work(self):
        """Демо - чтобы нажимать: обычные формы пишут как всегда."""
        r = self.client.post("/bikes", data={"code": "D-1", "model": "Kugoo"})
        self.assertEqual(r.status_code, 303)
        self.assertIsNotNone(run(self.crm.bike_by_code("D-1")))


class TestCheckPhoto(DemoCase):
    def setUp(self):
        super().setUp()
        self.assertEqual(self.login().status_code, 303)

    def test_bike_check_passes_but_photo_is_not_saved(self):
        bike_id = run(self.crm.create_bike(code="D-2", model="Kugoo", frame_no="F-2",
                                           status="new"))
        r = self.client.post(f"/bikes/{bike_id}/check",
                             data={"action": "check", "field": "frame_no"},
                             files={"photo": ("рама.jpg", b"\xff\xd8\xff" + b"0" * 64,
                                              "image/jpeg")})
        self.assertEqual(r.status_code, 303)
        mark = run(self.crm.bike(bike_id))["checked"]["frame_no"]
        self.assertFalse(mark.get("photo"))
        self.assertFalse(Path(self.cfg.bike_photo_dir).exists())
        text = self.get_ok(f"/bikes/{bike_id}")
        self.assertIn(web_app.DEMO_PHOTO_TEXT, text)

    def test_battery_check_passes_but_photo_is_not_saved(self):
        battery_id = run(self.crm.create_battery(code="9510009", serial_no="SN-9",
                                                 status="new"))
        r = self.client.post(f"/batteries/{battery_id}/check",
                             data={"action": "check", "field": "serial_no"},
                             files={"photo": ("корпус.png", b"\x89PNG" + b"0" * 64,
                                              "image/png")})
        self.assertEqual(r.status_code, 303)
        mark = run(self.crm.battery(battery_id))["checked"]["serial_no"]
        self.assertFalse(mark.get("photo"))
        self.assertFalse(Path(self.cfg.bike_photo_dir).exists())


class TestPhotoRequirementIsOff(DemoCase):
    """«Требовать фото номера» в демо не включить: снимки не хранятся, и
    включённое требование заперло бы ввод техники всем посетителям - а
    оператор и механик выключить его не могут."""

    def setUp(self):
        super().setUp()
        self.assertEqual(self.login().status_code, 303)

    def test_intake_keeps_photo_off(self):
        text = self.get_ok("/intake")
        self.assertRegex(text, r'name="photo"\s*disabled')
        self.assertIn("это требование выключено", text)
        r = self.client.post("/intake", data={"required": "1", "photo": "1",
                                              "search_after_days": "7",
                                              "theft_after_days": "21"})
        self.assertEqual(r.status_code, 303)
        self.assertEqual(self.crm.settings_.get("bike_photo_required"), "0")
        self.assertEqual(self.crm.settings_.get("bike_check_required"), "1")
        self.assertIn(web_app.DEMO_NO_PHOTO_TEXT, self.get_ok("/intake"))

    def test_new_bike_can_still_be_commissioned(self):
        self.client.post("/intake", data={"required": "1", "photo": "1"})
        bike_id = run(self.crm.create_bike(code="D-3", model="Kugoo", frame_no="F-3",
                                           status="new"))
        text = self.get_ok(f"/bikes/{bike_id}")
        self.assertNotIn('type="file" name="photo"', text, "фото не просят")
        mechanic = self.client_from("10.0.9.1")
        self.assertEqual(self.login("mechanic", client=mechanic).status_code, 303)
        r = mechanic.post(f"/bikes/{bike_id}/check",
                          data={"action": "check", "field": "frame_no"})
        self.assertEqual(r.status_code, 303)
        self.assertIn("frame_no", run(self.crm.bike(bike_id))["checked"])


class TestOurTemplatesAreHidden(DemoCase):
    """Наши шаблоны договора и актов - с реквизитами настоящего ИП (ФИО,
    ИНН, счёт, телефоны). В публичном демо их не отдать и не показать."""

    def test_no_link_and_404(self):
        self.assertEqual(self.login().status_code, 303)
        text = self.get_ok("/documents")
        self.assertNotIn("/documents/ours/", text)
        self.assertNotIn("Скачать наш", text)
        for kind in ("contract", "act_in", "act_out", "buyout", "consent"):
            r = self.client.get(f"/documents/ours/{kind}")
            self.assertEqual(r.status_code, 404, kind)
            self.assertNotIn("attachment", r.headers.get("content-disposition", ""), kind)


class TestBodyLimit(DemoCase):
    """Тело запроса больше предела - 413 до разбора формы: иначе чужая
    загрузка на /login (он открыт без входа) заполняла бы диск сервера
    временными файлами ещё до проверки пароля."""

    def test_big_upload_to_login_is_refused_before_parsing(self):
        with mock.patch.object(web_app.logic, "check_login") as parsed:
            r = self.client.post("/login", data={"login": "demo", "password": "demo"},
                                 files={"x": ("big.bin", b"0" * (web_app.DEMO_BODY_MAX + 1),
                                              "application/octet-stream")})
            parsed.assert_not_called()
        self.assertEqual(r.status_code, 413)
        self.assertIn("Слишком большой запрос", r.text)

    def test_chunked_body_is_counted(self):
        """Без Content-Length тело считается по мере чтения."""
        def chunks():
            for _ in range(web_app.DEMO_BODY_MAX // 65536 + 2):
                yield b"a=" + b"0" * 65534
        r = self.client.post("/login", content=chunks(),
                             headers={"content-type": "application/x-www-form-urlencoded"})
        self.assertEqual(r.status_code, 413)

    def test_normal_forms_pass(self):
        self.assertEqual(self.login().status_code, 303)
        self.assertEqual(self.login(password="nope",
                                    client=self.client_from("10.0.8.1")).status_code, 401)


class TestRequestPace(DemoCase):
    """Предел запросов с одного адреса: цикл curl не замораживает демо
    остальным посетителям."""

    def test_limits_by_address(self):
        now = [0.0]
        limits = web_app.DemoLimits(inflight=2, rate=1.0, burst=3, clock=lambda: now[0])
        self.assertTrue(limits.enter("a"))
        self.assertTrue(limits.enter("a"))
        self.assertFalse(limits.enter("a"), "двое уже в работе")
        self.assertTrue(limits.enter("b"), "у другого адреса свой счёт")
        limits.leave("a")
        self.assertTrue(limits.enter("a"), "третий жетон ведра")
        limits.leave("a")
        limits.leave("a")
        self.assertFalse(limits.enter("a"), "ведро пусто")
        now[0] += 1.0
        self.assertTrue(limits.enter("a"), "жетон за секунду")

    def test_gate_answers_429_and_spares_static(self):
        self.app.state.demo_limits = web_app.DemoLimits(rate=0.0, burst=2)
        self.assertEqual(self.client.get("/login").status_code, 200)
        self.assertEqual(self.client.get("/login").status_code, 200)
        r = self.client.get("/login")
        self.assertEqual(r.status_code, 429)
        self.assertIn(web_app.DEMO_BUSY_TEXT, r.text)
        self.assertEqual(r.headers.get("x-robots-tag"), "noindex, nofollow")
        for path in ("/static/style.css", "/healthz"):
            self.assertEqual(self.client.get(path).status_code, 200, path)
        other = self.client_from("10.0.7.1")
        self.assertEqual(other.get("/login").status_code, 200)

    def test_every_demo_panel_has_limits(self):
        self.assertIsInstance(self.app.state.demo_limits, web_app.DemoLimits)


class TestExportsAreMarked(DemoCase):
    """Скачанная выгрузка демо помечена сама: имя файла, лист и первая
    строка. Файл пересылают без страницы с плашкой."""

    def setUp(self):
        super().setUp()
        self.assertEqual(self.login().status_code, 303)
        for i in range(3):
            run(self.crm.create_client(full_name=f"Клиент Демо {i}",
                                       phone=f"+7000000000{i}"))

    def test_xlsx(self):
        from openpyxl import load_workbook
        r = self.client.get("/clients.xlsx")
        self.assertEqual(r.status_code, 200)
        self.assertIn('filename="demo-clients.xlsx"', r.headers["content-disposition"])
        book = load_workbook(io.BytesIO(r.content))
        sheet = book.active
        self.assertEqual(sheet.title, "Демо")
        self.assertEqual(sheet["A1"].value, web_app.DEMO_EXPORT_NOTE)
        self.assertEqual(sheet.freeze_panes, "A3")
        names = [row[0] for row in sheet.iter_rows(min_row=3, values_only=True)]
        self.assertEqual(len(names), 3)
        self.assertEqual(book.properties.title, web_app.DEMO_EXPORT_NOTE)

    def test_csv(self):
        r = self.client.get("/clients.csv")
        self.assertIn('filename="demo-clients.csv"', r.headers["content-disposition"])
        lines = r.content.decode("utf-8-sig").splitlines()
        self.assertEqual(lines[0], web_app.DEMO_EXPORT_NOTE)
        self.assertEqual(len(lines), 2 + 3)

    def test_rows_are_capped_and_say_so(self):
        with mock.patch.object(web_app, "DEMO_EXPORT_ROWS", 2):
            lines = self.client.get("/clients.csv").content.decode("utf-8-sig").splitlines()
        self.assertEqual(lines[0], f"{web_app.DEMO_EXPORT_NOTE} Показаны первые 2 строк.")
        self.assertEqual(len(lines), 2 + 2)

    def test_file_names(self):
        self.assertEqual(web_app.demo_disposition('attachment; filename="a.csv"'),
                         'attachment; filename="demo-a.csv"')
        self.assertEqual(web_app.demo_disposition("attachment; filename*=utf-8''%D0%B0.docx"),
                         "attachment; filename*=utf-8''demo-%D0%B0.docx")
        self.assertEqual(web_app.demo_disposition('attachment; filename="demo-a.csv"'),
                         'attachment; filename="demo-a.csv"')


class TestLoginThrottle(DemoCase):
    def test_shared_login_is_not_locked(self):
        """Двенадцать чужих ошибок под demo с разных адресов не запирают
        вход остальным посетителям."""
        for i in range(web_app.LOGIN_LIMIT + 2):
            r = self.login(password="nope", client=self.client_from(f"10.0.2.{i}"))
            self.assertEqual(r.status_code, 401)
        self.assertEqual(self.login(client=self.client_from("10.0.3.1")).status_code, 303)

    def test_address_limit_still_holds(self):
        with mock.patch.object(web_app, "LOGIN_IP_LIMIT", 5):
            for _ in range(5):
                self.assertEqual(self.login(password="nope").status_code, 401)
            r = self.login()
            self.assertEqual(r.status_code, 429)
            self.assertIn("Слишком много попыток", r.text)
            # с другого адреса - пускает
            self.assertEqual(self.login(client=self.client_from("10.0.4.1")).status_code,
                             303)


class TestRobotsAndMaintenance(DemoCase):
    def test_noindex_on_every_response(self):
        for path in ("/login", "/healthz", "/static/style.css", "/", "/robots.txt"):
            r = self.client.get(path)
            self.assertEqual(r.headers.get("x-robots-tag"), "noindex, nofollow", path)
        self.login()
        self.assertEqual(self.client.get("/bikes").headers.get("x-robots-tag"),
                         "noindex, nofollow")

    def test_robots_txt_is_public(self):
        r = self.client.get("/robots.txt")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.text, "User-agent: *\nDisallow: /\n")
        self.assertTrue(r.headers["content-type"].startswith("text/plain"))

    def test_503_while_reset_runs(self):
        self.login()
        self.app.state.maintenance = True
        for method, path in (("GET", "/"), ("GET", "/login"), ("GET", "/static/style.css"),
                             ("POST", "/login"), ("GET", "/robots.txt")):
            r = self.client.request(method, path)
            self.assertEqual(r.status_code, 503, path)
            self.assertIn("Демо обновляется, минуту", r.text)
            self.assertEqual(r.headers.get("retry-after"), "60")
            self.assertEqual(r.headers.get("x-robots-tag"), "noindex, nofollow")
        self.assertEqual(self.client.get("/healthz").status_code, 200)
        self.app.state.maintenance = False
        self.get_ok("/")


class TestNormalModeUnchanged(DemoCase):
    demo = False

    def test_no_banner_no_hints_no_robots(self):
        text = self.get_ok("/login")
        self.assertNotIn(BANNER, text)
        self.assertNotIn("demo-logins", text)
        self.assertIn("Доступ выдаёт администратор", text)
        r = self.client.get("/login")
        self.assertNotIn("x-robots-tag", r.headers)
        # robots.txt в обычной панели нет - это закрытая страница
        self.assertEqual(self.client.get("/robots.txt").status_code, 303)
        self.assertIn("crm_session=", self.login().headers.get("set-cookie", ""))
        self.assertNotIn(BANNER, self.get_ok("/"))

    def test_exports_are_plain(self):
        self.assertEqual(self.login().status_code, 303)
        run(self.crm.create_client(full_name="Клиент Боевой", phone="+70000000009"))
        r = self.client.get("/clients.csv")
        self.assertIn('filename="clients.csv"', r.headers["content-disposition"])
        self.assertNotIn("Демо", r.content.decode("utf-8-sig"))
        self.assertIsNone(self.app.state.demo_limits)

    def test_import_and_our_templates_stay(self):
        self.assertEqual(self.login().status_code, 303)
        self.assertIn('action="/import"', self.get_ok("/import"))
        self.assertIn("/documents/ours/contract", self.get_ok("/documents"))
        r = self.client.get("/documents/ours/contract")
        self.assertEqual(r.status_code, 200)

    def test_body_limit_fits_the_import(self):
        """Боевой предел тела пропускает импорт на 20 МБ и не больше."""
        self.assertGreater(web_app.BODY_MAX, web_app.IMPORT_MAX_BYTES + 64 * 1024)
        with mock.patch.object(web_app, "BODY_MAX", 4096):
            app = create_app(crm=self.crm, db=FakeBotDB(), cfg=self.cfg, bot=None)
        client = TestClient(app, follow_redirects=False)
        r = client.post("/login", data={"login": "demo", "password": "x" * 5000})
        self.assertEqual(r.status_code, 413)
        r = client.post("/login", data={"login": "demo", "password": "demo"})
        self.assertEqual(r.status_code, 303)

    def test_maintenance_flag_is_ignored(self):
        self.app.state.maintenance = True
        self.get_ok("/login")

    def test_login_lock_is_per_login(self):
        for i in range(web_app.LOGIN_LIMIT):
            self.login(password="nope", client=self.client_from(f"10.0.5.{i}"))
        r = self.login(client=self.client_from("10.0.6.1"))
        self.assertEqual(r.status_code, 429)

    def test_password_and_staff_forms_work(self):
        self.assertEqual(self.login().status_code, 303)
        manager = run(self.crm.access_profile_by_code("manager"))
        count = len(self.crm.staff)
        self.client.post("/staff", data={"login": "newbie", "password": "newbie-123",
                                         "name": "Новичок", "profile_id": str(manager["id"])})
        self.assertEqual(len(self.crm.staff), count + 1)
        self.client.post("/me/password", data={"old": "demo", "new": "new-pass-123"})
        self.assertTrue(logic.verify_password("new-pass-123", self.hash_of("demo")))


# ─────────────────────────── процесс app.demo ───────────────────────────

@unittest.skipUnless(HAVE_DEMO, "нет пакета демо или asyncpg")
class TestDemoConfig(unittest.TestCase):
    ENV = {"POSTGRES_PASSWORD": "pg-pass", "POSTGRES_DB": "mybike_demo",
           "CRM_SECRET": "demo-cookie-secret", "BOT_TOKEN": "123456:REAL-TOKEN",
           "TOCHKA_TOKEN": "bank-token", "TOCHKA_CUSTOMER_CODE": "300000000",
           "INBOX_KEY": "A" * 44, "INBOX_HOOK_TOKEN": "hook-token",
           "PAY_URL": "https://qr.nspk.ru/real", "CONTRACT_CHAT_ID": "-100123",
           "CRM_TITLE": "МАЙБАЙК", "BIKE_PHOTO_DIR": "/bikes",
           "DOC_TEMPLATE_DIR": "/doctemplates", "STORAGE_DIR": "/files/kyc"}

    def load(self, **extra):
        with mock.patch.dict(os.environ, {**self.ENV, **extra}, clear=True):
            return runtime.load_config()

    def test_outside_keys_are_forced_empty(self):
        cfg = self.load()
        self.assertTrue(cfg.demo)
        for name in ("bot_token", "tochka_token", "tochka_customer_code",
                     "inbox_hook_token", "pay_url", "contract_chat_id", "admin_password"):
            self.assertEqual(getattr(cfg, name), "", name)
        self.assertEqual(cfg.title, "МАЙБАЙК · демо")
        self.assertEqual(cfg.secret, "demo-cookie-secret")
        self.assertEqual(cfg.pg["database"], "mybike_demo")

    def test_token_file_in_environment_is_ignored_too(self):
        with tempfile.NamedTemporaryFile("w", suffix=".txt", delete=False) as f:
            f.write("999:FROM-FILE")
        self.addCleanup(os.unlink, f.name)
        cfg = self.load(BOT_TOKEN_FILE=f.name)
        self.assertEqual(cfg.bot_token, "")

    def test_files_stay_in_tmp(self):
        cfg = self.load()
        for path in (cfg.bike_photo_dir, cfg.doc_dir, cfg.storage_dir):
            self.assertTrue(str(path).startswith("/tmp/"), path)

    def test_inbox_key_is_derived_and_reads_the_seed(self):
        cfg = self.load()
        self.assertNotEqual(cfg.inbox_key, self.ENV["INBOX_KEY"])
        Vault.from_raw(cfg.inbox_key)                   # формат, который ждёт панель
        # Сид шифрует тем же производным ключом, что получает панель.
        sealed = service.inbox_seal(service.inbox_vault(seed_extras.inbox_key_text(cfg.secret)),
                                    "Добрый день, велосипед ещё свободен?")
        opened = service.inbox_vault(cfg.inbox_key).decrypt(sealed)
        self.assertEqual(opened, {"t": "Добрый день, велосипед ещё свободен?"})

    def test_panel_from_demo_config_is_in_demo_mode(self):
        cfg = self.load()
        app = create_app(crm=FakeCrm(), db=FakeBotDB(), cfg=cfg, bot=None)
        client = TestClient(app, follow_redirects=False)
        r = client.get("/login")
        self.assertIn(BANNER, r.text)
        self.assertIn("МАЙБАЙК · демо", r.text)
        self.assertEqual(r.headers.get("x-robots-tag"), "noindex, nofollow")


@unittest.skipUnless(HAVE_DEMO, "нет пакета демо или asyncpg")
class TestNightlyReset(unittest.TestCase):
    def cfg(self):
        return WebConfig(pg={"database": "mybike_demo"}, secret="s", admin_login="admin",
                         admin_password="", bot_token="", storage_dir=Path("/tmp/kyc"),
                         port=8080, remind_before_days=2, demo=True)

    def app(self):
        state = type("State", (), {"maintenance": False})()
        return type("App", (), {"state": state})()

    def test_next_reset_is_four_am_moscow(self):
        at = runtime.next_reset
        self.assertEqual(at(datetime(2026, 9, 25, 3, 59, tzinfo=MSK)),
                         datetime(2026, 9, 25, 4, 0, tzinfo=MSK))
        self.assertEqual(at(datetime(2026, 9, 25, 4, 0, tzinfo=MSK)),
                         datetime(2026, 9, 26, 4, 0, tzinfo=MSK))
        self.assertEqual(at(datetime(2026, 9, 25, 12, 0, tzinfo=MSK)),
                         datetime(2026, 9, 26, 4, 0, tzinfo=MSK))
        # 00:30 UTC - это 03:30 в Москве, сброс в тот же день
        self.assertEqual(at(datetime(2026, 9, 25, 0, 30, tzinfo=UTC)),
                         datetime(2026, 9, 25, 4, 0, tzinfo=MSK))

    def test_reset_runs_under_maintenance_once_a_night(self):
        app, calls, sleeps = self.app(), [], []
        clock = [datetime(2026, 9, 25, 12, 0, tzinfo=MSK)]

        async def fake_reset(pool, *, today, now, secret, **kwargs):
            # Сид шифрует «Входящие» ключом от CRM_SECRET - тем же, что панель.
            self.assertEqual(secret, "s")
            calls.append((today, now, app.state.maintenance))
            return {"today": today.isoformat()}

        async def fake_sleep(seconds):
            sleeps.append(seconds)
            if len(sleeps) == 3:
                raise asyncio.CancelledError
            # просыпаемся на долю секунды раньше срока - так бывает
            clock[0] += timedelta(seconds=seconds - 0.2)

        with mock.patch.object(runtime.seed, "reset", fake_reset), \
                self.assertRaises(asyncio.CancelledError):
            asyncio.run(runtime.nightly(app, None, self.cfg(), clock=lambda: clock[0],
                                        sleep=fake_sleep))
        self.assertEqual(sleeps[0], 16 * 3600)
        self.assertAlmostEqual(sleeps[1], 24 * 3600, delta=1)
        self.assertEqual([c[0] for c in calls], [date(2026, 9, 26), date(2026, 9, 27)])
        self.assertTrue(all(c[2] for c in calls), "сброс - только под 503")
        self.assertFalse(app.state.maintenance)

    def test_failed_reset_is_retried_until_morning(self):
        """Один сбой ночью (база на миг недоступна) не оставляет демо
        вчерашним на весь день: повтор через RETRY_AFTER. Живой день до
        сброса останавливается, после удачного - начинается заново."""
        app, calls, sleeps, events = self.app(), [], [], []
        clock = [datetime(2026, 9, 30, 3, 0, tzinfo=MSK)]

        async def flaky(pool, *, today, now, secret, **kwargs):
            calls.append(now)
            if len(calls) == 1:
                raise ConnectionRefusedError("postgres-demo перезапускается")
            return {"today": today.isoformat(), "later": [{"kind": "shift"}]}

        async def fake_sleep(seconds):
            sleeps.append(seconds)
            if len(sleeps) == 3:
                raise asyncio.CancelledError
            clock[0] += timedelta(seconds=seconds)

        class Day:
            def stop(self):
                events.append("stop")

            def start(self, summary):
                events.append(("start", len(summary["later"])))

        with mock.patch.object(runtime.seed, "reset", flaky), \
                self.assertLogs("crm.demo", "WARNING"), \
                self.assertRaises(asyncio.CancelledError):
            asyncio.run(runtime.nightly(app, None, self.cfg(), clock=lambda: clock[0],
                                        sleep=fake_sleep, day=Day()))
        self.assertEqual(sleeps[0], 3600)
        self.assertEqual(sleeps[1], runtime.RETRY_AFTER.total_seconds())
        self.assertAlmostEqual(sleeps[2], 24 * 3600 - runtime.RETRY_AFTER.total_seconds(),
                               delta=1)
        self.assertEqual([c.strftime("%d %H:%M") for c in calls], ["30 04:00", "30 04:10"])
        self.assertEqual(events, ["stop", "stop", ("start", 1)])

    def test_retries_stop_in_the_morning(self):
        """После RETRY_UNTIL демо уже смотрят: сброс посреди показа хуже
        вчерашних чисел - следующая попытка ночью."""
        night = datetime(2026, 9, 30, 4, 0, tzinfo=MSK)
        self.assertEqual(runtime.after_failure(night, night),
                         night + runtime.RETRY_AFTER)
        late = datetime(2026, 9, 30, 6, 55, tzinfo=MSK)
        self.assertEqual(runtime.after_failure(late, night),
                         datetime(2026, 10, 1, 4, 0, tzinfo=MSK))
        # Сбой при запуске днём - тоже до следующей ночи, а не через 10 минут.
        day = datetime(2026, 9, 30, 15, 0, tzinfo=MSK)
        self.assertEqual(runtime.after_failure(day, day),
                         datetime(2026, 10, 1, 4, 0, tzinfo=MSK))

    def test_failed_reset_keeps_yesterday_and_reopens(self):
        app = self.app()

        async def broken(pool, **kwargs):
            raise RuntimeError("сид упал")

        with mock.patch.object(runtime.seed, "reset", broken), \
                self.assertLogs("crm.demo", "ERROR") as logs:
            ok = asyncio.run(runtime.reset_in_maintenance(app, None, self.cfg()))
        self.assertFalse(ok)
        self.assertFalse(app.state.maintenance, "панель снова отвечает - вчерашним демо")
        self.assertIn("прежние данные", logs.output[0])


class FakePool:
    """fetchval по началу запроса: какая проверка - такой ответ."""

    def __init__(self, *, schema: bool, foreign: bool = False, demo: bool = True):
        self.answers = {"select to_regclass": not schema,
                        "select exists (select 1 from crm.staff) and not": foreign,
                        "select exists (select 1 from crm.staff where": demo}

    async def fetchval(self, sql):
        return next(v for k, v in self.answers.items() if sql.startswith(k))


@unittest.skipUnless(HAVE_DEMO, "нет пакета демо или asyncpg")
class TestDatabaseGuard(unittest.TestCase):
    def test_refuses_a_database_not_named_demo(self):
        with self.assertRaisesRegex(RuntimeError, "не похожа на демо"):
            asyncio.run(runtime.check_database(FakePool(schema=False), "mybike"))

    def test_refuses_foreign_staff(self):
        with self.assertRaisesRegex(RuntimeError, "нет логина demo"):
            asyncio.run(runtime.check_database(FakePool(schema=True, foreign=True),
                                               "mybike_demo"))

    def test_accepts_empty_or_own_database(self):
        asyncio.run(runtime.check_database(FakePool(schema=False), "mybike_demo"))
        asyncio.run(runtime.check_database(FakePool(schema=True), "mybike_demo"))

    def test_has_demo(self):
        self.assertFalse(asyncio.run(runtime.has_demo(FakePool(schema=False))))
        self.assertTrue(asyncio.run(runtime.has_demo(FakePool(schema=True))))
        self.assertFalse(asyncio.run(runtime.has_demo(FakePool(schema=True, demo=False))))


if __name__ == "__main__":                              # pragma: no cover
    unittest.main()
