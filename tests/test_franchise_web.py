"""Франшиза в панели: адрес /hook/metrics у франчайзи и раздел «Франчайзи»
у франчайзера - права, карточка, чужие строки на экране, кнопка «Обновить
сейчас», роялти и выгрузка. База - tests/fake_crm.py, сеть - заглушка.
"""

from __future__ import annotations

import dataclasses
import io
import shutil
import sys
import unittest
from datetime import datetime, timedelta
from decimal import Decimal
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

try:
    from app.crm import logic, service
    from app.services.crypto import generate_key
    from app.web import app as web_app
    from tests import test_web as tw
    from tests.test_franchise import PII_KEYS, keys_of, payload, raw
    HAVE_WEB = tw.HAVE_WEB
except ImportError:                                     # pragma: no cover
    HAVE_WEB = False

D = Decimal
TOKEN = "metrics-" + "a" * 40


@unittest.skipUnless(HAVE_WEB, "fastapi не установлен")
class TestMetricsHook(tw.WebCase):
    """Сторона франчайзи: только агрегаты, только по токену, не чаще предела."""

    def build(self, **over):
        self.cfg = dataclasses.replace(self.cfg, **over)
        self.app = tw.create_app(crm=self.crm, db=self.db, cfg=self.cfg, bot=self.bot)
        self.client = tw.TestClient(self.app, follow_redirects=False)

    def get(self, token=TOKEN):
        headers = {"Authorization": f"Bearer {token}"} if token is not None else {}
        return self.client.get("/hook/metrics", headers=headers)

    def test_off_without_token_and_in_demo(self):
        r = self.get()
        self.assertEqual(r.status_code, 404)
        self.assertNotIn("location", r.headers, "не редирект на вход")
        self.build(metrics_token=TOKEN, demo=True)
        self.assertEqual(self.get().status_code, 404, "у демо метрик наружу нет")

    def test_token_is_checked(self):
        self.build(metrics_token=TOKEN)
        self.assertEqual(self.get(None).status_code, 401)
        self.assertEqual(self.get("x" + TOKEN[1:]).status_code, 401)
        self.assertEqual(self.client.get("/hook/metrics", headers={
            "Authorization": TOKEN}).status_code, 401, "только Bearer")
        r = self.get()
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.headers["cache-control"], "no-store")
        self.assertNotIn("set-cookie", r.headers, "адрес сессию не заводит")
        self.assertTrue(r.headers["content-type"].startswith("application/json"))

    def test_payload_has_aggregates_only(self):
        self.seed()
        tw.run(self.crm.create_location(name="Павлюхина", city="Казань",
                                        address="ул. Тайная, 5", note=None))
        tw.run(self.crm.add_ledger(client_id=self.client_id, kind="payment",
                                   amount=D("3000"), note="Иванов, паспорт 9200"))
        self.build(metrics_token=TOKEN)
        r = self.get()
        data = r.json()
        self.assertEqual(keys_of(data) & PII_KEYS, set())
        for secret in ("Иванов", "+79990000000", "5001", "Тайная", "паспорт"):
            self.assertNotIn(secret, r.text)
        self.assertEqual(data["clients"], 1)
        self.assertEqual(data["name"], "МАЙБАЙК")
        self.assertEqual(data["city"], "Казань")
        self.assertEqual(len(data["version"]), 12)
        self.assertTrue(logic.parse_metrics(data).ok, "франчайзер примет свой же ответ")

    def test_wrong_tokens_are_throttled(self):
        self.build(metrics_token=TOKEN)
        for _ in range(logic.HOOK_FAIL_LIMIT):
            self.assertEqual(self.get("wrong").status_code, 401)
        self.assertEqual(self.get("wrong").status_code, 429)

    def test_valid_requests_are_rate_limited(self):
        self.build(metrics_token=TOKEN)
        with mock.patch.object(logic, "METRICS_RATE_LIMIT", 3):
            self.assertEqual([self.get().status_code for _ in range(4)],
                             [200, 200, 200, 429])

    def test_hook_is_public_and_outside_sections(self):
        self.assertTrue("/hook/metrics".startswith(web_app.PUBLIC))
        self.assertIsNone(logic.section_for("/hook/metrics"))
        self.assertEqual(logic.section_for("/franchisees/royalty.csv"), "franchise")
        self.assertEqual(logic.section_for("/franchisees/refresh/1"), "franchise")


@unittest.skipUnless(HAVE_WEB, "fastapi не установлен")
class TestFranchiseSection(tw.WebCase):
    """Сторона франчайзера: раздел только у владельца, токен шифром, чужие
    строки - экранированными."""

    def setUp(self):
        super().setUp()
        self.cfg = dataclasses.replace(self.cfg, franchise_key=generate_key())
        self.app = tw.create_app(crm=self.crm, db=self.db, cfg=self.cfg, bot=self.bot)
        self.client = tw.TestClient(self.app, follow_redirects=False)
        self.vault = service.franchise_vault(self.cfg.franchise_key)
        self.login()

    def form(self, **over):
        data = {"name": "Май Байк Самара", "city": "Самара",
                "base_url": "https://crm.samara.example", "token": "t" * 32,
                "royalty_percent": "5", "fixed_fee": "15000",
                "contract_start": "2026-01-10", "active": "1"}
        data.update(over)
        return data

    def create(self, **over):
        r = self.client.post("/franchisees/new", data=self.form(**over))
        self.assertEqual(r.status_code, 303, r.text[:500])
        return int(r.headers["location"].rsplit("/", 1)[1])

    def store(self, fid, **over):
        now = datetime.now().astimezone()
        cur = now.date().replace(day=1)
        prev = (cur - timedelta(days=1)).replace(day=1)
        body = payload(generated_at=now.isoformat(timespec="seconds"),
                       months=[{"month": prev.strftime("%Y-%m"), "partial": False,
                                "idle_percent": 1, "avg_check": "1", "revenue": "100000",
                                "operational_days": "100", "rented_days": "90",
                                "idle_days": "10"}], **over)
        parsed = logic.parse_metrics_bytes(raw(body), now=now)
        self.assertTrue(parsed.ok, parsed.error)
        tw.run(service.franchise_store(self.crm, fid, parsed.value, today=now.date()))

    def test_only_owner_sees_the_section(self):
        self.assertIn(">Франчайзи</a>", self.get_ok("/"))
        manager = tw.run(self.crm.access_profile_by_code("manager"))
        tw.run(self.crm.create_staff("anna", logic.hash_password("password-1"), "Анна",
                                     "manager", manager["id"]))
        self.client.post("/logout")
        self.login("anna", "password-1")
        self.assertNotIn(">Франчайзи</a>", self.get_ok("/"))
        for path in ("/franchisees", "/franchisees/royalty", "/franchisees/royalty.csv",
                     "/franchisees/new"):
            self.assertEqual(self.client.get(path).status_code, 403, path)
        self.assertEqual(self.client.post("/franchisees/new", data=self.form())
                         .status_code, 403)
        self.assertEqual(self.crm.franchisees_, {})

    def test_create_edit_and_token(self):
        fid = self.create()
        row = tw.run(self.crm.franchisee(fid))
        self.assertNotIn("t" * 32, row["token_enc"], "токен - только шифром")
        self.assertEqual(service.franchise_token(self.vault, row), "t" * 32)
        page = self.get_ok(f"/franchisees/{fid}")
        self.assertNotIn("t" * 32, page, "токен на странице не показывается")
        self.assertIn("Ответа от франчайзи ещё не было", page)
        # Правка без токена - прежний токен остаётся.
        r = self.client.post(f"/franchisees/{fid}", data=self.form(token="", fixed_fee="0"))
        self.assertEqual(r.status_code, 303)
        row = tw.run(self.crm.franchisee(fid))
        self.assertEqual(service.franchise_token(self.vault, row), "t" * 32)
        self.assertEqual(row["fixed_fee"], D("0.00"))
        self.assertIn("Май Байк Самара", self.get_ok("/franchisees"))
        # Ошибка правки - та же форма с введённым, токен не затронут.
        r = self.client.post(f"/franchisees/{fid}",
                             data=self.form(token="", royalty_percent="300", city="Тольятти"))
        self.assertEqual(r.status_code, 400)
        self.assertIn('value="Тольятти"', r.text)
        self.assertNotIn("Ответа от франчайзи ещё не было", r.text)
        self.assertEqual(tw.run(self.crm.franchisee(fid))["city"], "Самара")

    def test_bad_forms_are_refused(self):
        for over in ({"base_url": "http://crm.samara.example"},
                     {"base_url": "javascript:alert(1)"}, {"royalty_percent": "120"},
                     {"contract_start": ""}, {"token": "short"}, {"token": ""}):
            r = self.client.post("/franchisees/new", data=self.form(**over))
            self.assertEqual(r.status_code, 400, over)
        self.assertEqual(self.crm.franchisees_, {})
        self.assertIn("Адрес: только https", self.client.post(
            "/franchisees/new", data=self.form(base_url="http://evil.example")).text)

    def test_no_key_no_token(self):
        self.cfg = dataclasses.replace(self.cfg, franchise_key="")
        self.app = tw.create_app(crm=self.crm, db=self.db, cfg=self.cfg, bot=self.bot)
        self.client = tw.TestClient(self.app, follow_redirects=False)
        self.login()
        self.assertIn("secrets/franchise_key", self.get_ok("/franchisees"))
        r = self.client.post("/franchisees/new", data=self.form())
        self.assertEqual(r.status_code, 400)
        self.assertEqual(self.crm.franchisees_, {})

    def test_hostile_strings_are_escaped(self):
        """Имя, город и точки из ответа франчайзи - текст, а не разметка."""
        fid = self.create()
        self.store(fid, name="<script>alert(1)</script>", city="<b>Город</b>",
                   version="<img src=x>",
                   points=[{"name": "<img src=x onerror=alert(1)>", "city": "\"'>"}])
        for path in (f"/franchisees/{fid}", "/franchisees"):
            page = self.get_ok(path)
            self.assertNotIn("<script>alert", page, path)
            self.assertNotIn("<img src=x", page, path)
            self.assertNotIn("<b>Город", page, path)
        page = self.get_ok(f"/franchisees/{fid}")
        self.assertIn("&lt;script&gt;alert(1)&lt;/script&gt;", page)
        self.assertIn("&lt;img src=x onerror=alert(1)&gt;", page)
        self.assertIn("не сообщена", page, "негодная версия отброшена")

    def test_comparison_and_royalty(self):
        fid = self.create()
        self.store(fid)
        page = self.get_ok("/franchisees")
        self.assertIn("свежие", page)
        # 100 000 × 5 % + 15 000 за прошлый месяц
        self.assertIn(tw.plain(logic.money(D("20000"))), page)
        self.assertIn("10.0 %", page, "простой пересчитан из дней: 10 из 100")
        royalty = self.get_ok("/franchisees/royalty")
        self.assertIn(tw.plain(logic.money(D("20000"))), royalty)
        self.get_ok("/franchisees/royalty?months=12")
        r = self.client.get("/franchisees/royalty.csv")
        self.assertEqual(r.status_code, 200)
        text = r.content.decode("utf-8-sig")
        self.assertIn("20000,00", text)
        self.assertIn("Май Байк Самара", text)
        r = self.client.get("/franchisees/royalty.xlsx")
        self.assertEqual(r.status_code, 200)
        from openpyxl import load_workbook
        sheet = load_workbook(io.BytesIO(r.content)).active
        values = [c.value for row in sheet.iter_rows() for c in row]
        self.assertIn(20000.0, values, "роялти - числом, а не строкой")
        self.assertEqual(self.client.get("/franchisees/royalty.pdf").status_code, 404)

    def test_stale_and_failed_are_marked(self):
        fid = self.create()
        tw.run(self.crm.franchise_failed(fid, "франчайзи не отвечает: TimeoutError"))
        page = self.get_ok("/franchisees")
        self.assertIn("нет свежих данных", page)
        self.assertIn("франчайзи не отвечает: TimeoutError", page)

    def test_refresh_button_fetches_once(self):
        fid = self.create()
        now = datetime.now().astimezone()
        calls = []

        async def fetch(url, token, **_):
            calls.append((url, token))
            return raw(payload(generated_at=now.isoformat(timespec="seconds")))

        with mock.patch("app.crm.franchise.franchise_http.fetch_metrics", fetch):
            r = self.client.post(f"/franchisees/refresh/{fid}")
        self.assertEqual(r.status_code, 303)
        self.assertEqual(calls, [("https://crm.samara.example/hook/metrics", "t" * 32)])
        row = tw.run(self.crm.franchisee(fid))
        self.assertIsNotNone(row["ok_at"])
        self.assertIn("Данные франчайзи обновлены", self.get_ok(f"/franchisees/{fid}"))

        async def broken(url, token, **_):
            return b'{"format": 1, "fleet": NaN}'

        with mock.patch("app.crm.franchise.franchise_http.fetch_metrics", broken):
            self.client.post(f"/franchisees/refresh/{fid}")
        page = self.get_ok(f"/franchisees/{fid}")
        self.assertIn("ответ не прошёл проверку", page)
        self.assertIn("64", page, "прежние цифры остались")

    def refresh_with(self, fid, body):
        async def fetch(url, token, **_):
            return body

        with mock.patch("app.crm.franchise.franchise_http.fetch_metrics", fetch):
            return self.client.post(f"/franchisees/refresh/{fid}")

    def test_hostile_answers_do_not_break_the_card(self):
        """Год 0001 с поясом +14:00 ронял карточку (OverflowError в шаблоне),
        а «1e99999999999999999999» - саму кнопку (InvalidOperation): 500 и
        причина нигде. Теперь - отказ с причиной, карточка открывается, и
        франчайзи можно выключить или поправить."""
        fid = self.create()
        cases = ((raw(payload(generated_at="0001-01-01T00:00:00+14:00", months=[])),
                  "generated_at"),
                 (b'{"format": 1, "fleet": 1e99999999999999999999}', "не JSON"))
        for body, fragment in cases:
            r = self.refresh_with(fid, body)
            self.assertEqual(r.status_code, 303, fragment)
            row = tw.run(self.crm.franchisee(fid))
            self.assertIn(fragment, row["error"])
            self.assertIsNone(row["ok_at"], "негодный ответ не принят")
            self.assertIn(fragment, self.get_ok(f"/franchisees/{fid}"))

    def test_unforeseen_parse_failure_is_a_reason_not_500(self):
        fid = self.create()
        with mock.patch.object(logic, "parse_metrics_bytes", side_effect=TypeError("x")), \
                self.assertLogs("app.crm.franchise", "ERROR"):
            r = self.refresh_with(fid, raw(payload()))
        self.assertEqual(r.status_code, 303)
        self.assertIn("сбой разбора (TypeError)", tw.run(self.crm.franchisee(fid))["error"])

    def test_long_foreign_names_wrap_on_phone(self):
        """Имя и точка из ответа до 80 знаков без пробела: абзац переносит
        их в любом месте, иначе карточка уезжала вбок на телефоне."""
        fid = self.create()
        self.store(fid, name="M" * 80, points=[{"name": "W" * 80, "city": None}])
        page = self.get_ok(f"/franchisees/{fid}")
        self.assertIn('<p class="muted small wrap">Прислано', page)
        self.assertIn('<h1 class="wrap">', page)
        css = (Path(web_app.__file__).parent / "static" / "style.css").read_text()
        self.assertIn(".wrap{overflow-wrap:anywhere}", css)

    def test_terms_can_be_fixed_backwards_only_explicitly(self):
        """Опечатка в проценте, замеченная после первого опроса, иначе
        навсегда оставалась в прошлом месяце, который сейчас идёт в счёт."""
        fid = self.create(royalty_percent="5.5")
        self.store(fid)
        prev = (datetime.now().date().replace(day=1) - timedelta(days=1)).replace(day=1)

        def prev_terms():
            return next((m["royalty_percent"], m["fixed_fee"])
                        for m in tw.run(self.crm.franchise_months(prev))
                        if m["month"] == prev)

        self.assertEqual(prev_terms(), (D("5.50"), D("15000.00")))
        r = self.client.post(f"/franchisees/{fid}", data=self.form(token="",
                                                                   royalty_percent="7"))
        self.assertEqual(r.status_code, 303)
        self.assertEqual(prev_terms(), (D("5.50"), D("15000.00")),
                         "без поля - прошлые месяцы не переписываются")
        r = self.client.post(f"/franchisees/{fid}", data=self.form(
            token="", royalty_percent="7", terms_from=prev.strftime("%Y-%m")))
        self.assertEqual(r.status_code, 303)
        self.assertEqual(prev_terms(), (D("7.00"), D("15000.00")))
        # 100 000 × 7 % + 15 000
        self.assertIn(tw.plain(logic.money(D("22000"))), self.get_ok("/franchisees/royalty"))
        future = (datetime.now().date().replace(day=1) + timedelta(days=32)).strftime("%Y-%m")
        r = self.client.post(f"/franchisees/{fid}", data=self.form(
            token="", royalty_percent="9", terms_from=future))
        self.assertEqual(r.status_code, 400)
        self.assertIn("Условия с месяца", r.text)
        self.assertEqual(tw.run(self.crm.franchisee(fid))["royalty_percent"], D("7.00"))

    def test_mistaken_twin_can_be_wiped(self):
        """Дубль после первого опроса держался историей и удваивал итог
        роялти; стереть его можно только явной галочкой."""
        fid = self.create()
        twin = self.create(name="Май Байк Самара (дубль)")
        self.store(fid)
        self.store(twin)
        self.assertIn(tw.plain(logic.money(D("40000"))), self.get_ok("/franchisees/royalty"))
        self.client.post(f"/franchisees/{twin}/delete")
        self.assertIn(twin, self.crm.franchisees_, "без галочки история держит")
        r = self.client.post(f"/franchisees/{twin}/delete", data={"wipe": "1"})
        self.assertEqual((r.status_code, r.headers["location"]), (303, "/franchisees"))
        self.assertNotIn(twin, self.crm.franchisees_)
        self.assertFalse(any(k[0] == twin for k in self.crm.franchise_months_))
        royalty = self.get_ok("/franchisees/royalty")
        self.assertNotIn("дубль", royalty)
        self.assertNotIn(tw.plain(logic.money(D("40000"))), royalty)
        self.assertIn(fid, self.crm.franchisees_)

    def test_delete_only_without_history(self):
        empty = self.create(name="Пустой")
        self.assertEqual(self.client.post(f"/franchisees/{empty}/delete").status_code, 303)
        self.assertNotIn(empty, self.crm.franchisees_)
        fid = self.create()
        self.store(fid)
        self.client.post(f"/franchisees/{fid}/delete")
        self.assertIn(fid, self.crm.franchisees_, "история роялти держит")
        for path in ("/franchisees/999", "/franchisees/abc", "/franchisees/²",
                     "/franchisees/99999999999999999999999"):
            self.assertEqual(self.client.get(path).status_code, 404, path)
        self.assertEqual(self.client.post("/franchisees/refresh/x").status_code, 404)

    def test_demo_refresh_is_blocked(self):
        self.assertTrue(web_app.demo_blocked("POST", "/franchisees/refresh/1"))
        self.assertFalse(web_app.demo_blocked("POST", "/franchisees/1"))
        self.assertFalse(web_app.demo_blocked("GET", "/franchisees"))
        fid = self.create()
        self.cfg = dataclasses.replace(self.cfg, demo=True)
        self.app = tw.create_app(crm=self.crm, db=self.db, cfg=self.cfg, bot=None)
        self.client = tw.TestClient(self.app, follow_redirects=False)
        self.app.state.demo_limits = None
        self.assertEqual(self.login().headers["location"], "/")

        async def fetch(url, token, **_):                # pragma: no cover
            raise AssertionError("в демо сети нет")

        with mock.patch("app.crm.franchise.franchise_http.fetch_metrics", fetch):
            r = self.client.post(f"/franchisees/refresh/{fid}",
                                 headers={"referer": f"http://testserver/franchisees/{fid}"})
        self.assertEqual((r.status_code, r.headers["location"]),
                         (303, f"/franchisees/{fid}"))
        self.assertIn(web_app.DEMO_BLOCKED_TEXT, self.get_ok(f"/franchisees/{fid}"))
        self.assertIsNone(tw.run(self.crm.franchisee(fid))["polled_at"])


@unittest.skipUnless(HAVE_WEB, "fastapi не установлен")
class TestSecretsWiring(unittest.TestCase):
    """Новые секреты доезжают до своих сервисов и создаются установкой."""

    ROOT = Path(__file__).resolve().parent.parent

    def test_bootstrap_creates_every_compose_secret(self):
        compose = (self.ROOT / "docker-compose.yml").read_text(encoding="utf-8")
        boot = (self.ROOT / "bootstrap.sh").read_text(encoding="utf-8")
        for name in ("franchise_key", "metrics_token"):
            self.assertIn(f"file: ./secrets/{name}", compose)
            self.assertIn(f"secrets/{name}", boot)
        self.assertIn("if [ ! -s secrets/franchise_key ]", boot, "ключ генерируется")
        self.assertIn("[ -f secrets/metrics_token ] || : > secrets/metrics_token", boot,
                      "токен - пустая заглушка: адрес выключен")

    def test_demo_config_has_no_real_franchise_secrets(self):
        try:
            from app.demo import runtime
        except ImportError:                             # pragma: no cover
            self.skipTest("нет пакета демо или asyncpg")
        from app.web.config import WebConfig
        real = generate_key()
        cfg = runtime.demo_config(WebConfig(
            pg={}, secret="demo-secret-x", admin_login="admin", admin_password="",
            bot_token="", storage_dir=Path("/tmp"), port=8080, remind_before_days=2,
            metrics_token=TOKEN, franchise_key=real))
        self.assertEqual(cfg.metrics_token, "", "метрик наружу у демо нет")
        self.assertNotEqual(cfg.franchise_key, real, "боевой ключ в демо не попадает")
        self.assertIsNotNone(service.franchise_vault(cfg.franchise_key))

    @unittest.skipUnless(shutil.which("bash") and shutil.which("openssl"),
                         "нужны bash и openssl")
    def test_bootstrap_key_is_a_vault_key(self):
        """Ключ, который генерирует bootstrap.sh, принимает Vault: иначе
        раздел «Франчайзи» молча не хранил бы токены."""
        import re
        import subprocess
        import tempfile
        boot = (self.ROOT / "bootstrap.sh").read_text(encoding="utf-8")
        m = re.search(r"if \[ ! -s secrets/franchise_key \]; then\n(.*?)\nfi\n", boot, re.S)
        self.assertIsNotNone(m)
        with tempfile.TemporaryDirectory() as tmp:
            r = subprocess.run(["bash", "-c", "set -euo pipefail; mkdir -p secrets; "
                                + m.group().replace("say ", "echo ")],
                               cwd=tmp, capture_output=True, text=True)
            self.assertEqual(r.returncode, 0, r.stderr)
            key = (Path(tmp) / "secrets" / "franchise_key").read_text()
        self.assertIsNotNone(service.franchise_vault(key))

    def test_configs_read_the_files(self):
        import os
        import tempfile

        from app.config import Config
        from app.web.config import WebConfig
        with tempfile.TemporaryDirectory() as tmp:
            key = Path(tmp) / "franchise_key"
            key.write_text(generate_key())
            token = Path(tmp) / "metrics_token"
            token.write_text(TOKEN + "\n")
            env = {"FRANCHISE_KEY_FILE": str(key), "METRICS_TOKEN_FILE": str(token),
                   "POSTGRES_PASSWORD": "x", "CRM_SECRET": "s"}
            with mock.patch.dict(os.environ, env):
                web = WebConfig.load()
            self.assertEqual(web.metrics_token, TOKEN)
            self.assertEqual(web.franchise_key, key.read_text())
        fields = {f.name for f in dataclasses.fields(Config)}
        self.assertIn("franchise_key", fields)
        self.assertNotIn("metrics_token", fields, "адрес метрик - дело панели")


if __name__ == "__main__":
    unittest.main()
