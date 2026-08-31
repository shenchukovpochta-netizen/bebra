"""Выгрузка одной группы из xlsx-расписания в читаемый вид.

Нужна, чтобы пересобрать timetable/data.py на новый семестр: файл
расписания приходит с объединёнными ячейками и свободным текстом внутри,
и разбирать его вслепую нельзя - сначала надо увидеть, что там написано.

    python3 tools/dump_xlsx.py raspisanie.xlsx "3 курс" 06-445

Требует openpyxl (pip install openpyxl); в зависимости бота он не входит -
боту xlsx не нужен, он работает с уже разобранным timetable/data.py.
"""

from __future__ import annotations

import sys


def main(argv: list[str]) -> int:
    if len(argv) != 4:
        print(__doc__)
        return 2
    path, sheet, group = argv[1], argv[2], argv[3]
    try:
        import openpyxl
    except ImportError:
        print("нужен openpyxl: pip install openpyxl", file=sys.stderr)
        return 1

    ws = openpyxl.load_workbook(path, data_only=True)[sheet]
    # Значение объединённой ячейки лежит только в её левом верхнем углу,
    # остальные читаются как None. Без этой карты пропадут и дни недели,
    # и время пар - они как раз объединены на всю высоту блока.
    anchor: dict[tuple[int, int], tuple[int, int]] = {}
    for rng in ws.merged_cells.ranges:
        for row in range(rng.min_row, rng.max_row + 1):
            for col in range(rng.min_col, rng.max_col + 1):
                anchor[(row, col)] = (rng.min_row, rng.min_col)

    def val(row: int, col: int):
        row, col = anchor.get((row, col), (row, col))
        return ws.cell(row, col).value

    column = next(
        (c.column for row in ws.iter_rows() for c in row
         if c.value is not None and str(c.value).replace(" ", "") == group.replace(" ", "")),
        None,
    )
    if column is None:
        print(f"группа {group} на листе {sheet!r} не найдена", file=sys.stderr)
        return 1
    print(f"# {sheet}, группа {group}: столбец {column}")

    seen: set[tuple] = set()
    for row in range(1, ws.max_row + 1):
        day, slot, cell = val(row, 1), val(row, 2), val(row, column)
        key = (anchor.get((row, 1)), anchor.get((row, 2)), anchor.get((row, column)))
        if key in seen or (day is None and slot is None and cell is None):
            continue
        seen.add(key)
        print(f"\n### строка {row} | день={day!r} | слот={slot!r}")
        if cell is not None:
            print(repr(cell))
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
